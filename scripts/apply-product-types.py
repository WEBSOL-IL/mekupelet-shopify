#!/usr/bin/env python3
"""Write `productType` + the Shopify standard category to every product from the reviewed sheet.

Input: docs/import/out/product-type-proposal.xlsx (sheet "product types": handle, title, vendor, type;
sheet "scheme": type -> standard category full name). Rows whose type is "לבדיקה" or empty are
skipped and listed. The standard category full name is resolved once per type through
`taxonomy.categories(search:)` (exact fullName match, else the longest matching prefix, else none).
Dry run by default: prints the plan and the unresolved categories. `--execute` runs productUpdate per
product (only where type or category differ) and writes a report to docs/import/out/.

Run: python3 -I scripts/apply-product-types.py docs/import/out/product-type-proposal.xlsx [--execute] [--limit N]
"""
import argparse, importlib.util, json, os, sys, time, collections

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('imp', os.path.join(HERE, 'import-collections.py'))
imp = importlib.util.module_from_spec(spec); spec.loader.exec_module(imp)


def gql(query, variables, tries=4):
    """imp.gql with retries: the sandbox proxy drops TLS connections now and then."""
    for i in range(tries):
        try:
            return imp.gql(query, variables)
        except Exception as e:
            if i == tries - 1:
                raise
            time.sleep(2 * (i + 1))


def resolve_category(full_name, cache):
    if not full_name:
        return None
    if full_name in cache:
        return cache[full_name]
    parts = [p.strip() for p in full_name.split('>')]
    found = None
    for depth in range(len(parts), 0, -1):
        name = ' > '.join(parts[:depth])
        r = gql('query($s: String!) { taxonomy { categories(first: 25, search: $s) { nodes { id fullName } } } }', {'s': parts[depth - 1]})
        nodes = r['data']['taxonomy']['categories']['nodes']
        hit = next((n for n in nodes if n['fullName'] == name), None)
        if hit:
            found = (hit['id'], hit['fullName'], depth == len(parts))
            break
    cache[full_name] = found
    return found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('sheet'); ap.add_argument('--execute', action='store_true'); ap.add_argument('--limit', type=int, default=0)
    a = ap.parse_args()
    from openpyxl import load_workbook
    wb = load_workbook(a.sheet)
    scheme = {r[0]: (r[2] or '') for r in wb['scheme'].iter_rows(min_row=2, values_only=True) if r[0]}
    rows = [r for r in wb['product types'].iter_rows(min_row=2, values_only=True) if r[0]]
    todo, skipped = [], []
    for handle, title, vendor, ptype, *_ in rows:
        ptype = (ptype or '').strip()
        if not ptype or ptype == 'לבדיקה':
            skipped.append((handle, title)); continue
        if ptype not in scheme:
            scheme[ptype] = ''  # a type the merchant added by hand: type only, no category
        todo.append((handle, title, ptype))
    print(f'rows {len(rows)}, to write {len(todo)}, skipped {len(skipped)}')
    cache = {}
    unresolved = []
    for t, full in scheme.items():
        res = resolve_category(full, cache)
        if full and not res:
            unresolved.append((t, full))
        elif res and not res[2]:
            print(f'  category fallback for {t}: "{full}" -> "{res[1]}"')
    if unresolved:
        print('no standard category found for:', unresolved)
    # current state (type + category) to skip unchanged products
    current = {}
    cursor = None
    while True:
        r = gql('query($c: String) { products(first: 250, after: $c) { pageInfo { hasNextPage endCursor } nodes { id handle productType category { id } } } }', {'c': cursor})
        d = r['data']['products']
        for n in d['nodes']:
            current[n['handle']] = (n['id'], n['productType'] or '', (n['category'] or {}).get('id'))
        if not d['pageInfo']['hasNextPage']:
            break
        cursor = d['pageInfo']['endCursor']
    plan = []
    for handle, title, ptype in todo:
        if handle not in current:
            skipped.append((handle, 'not in store')); continue
        gid, cur_type, cur_cat = current[handle]
        res = cache.get(scheme.get(ptype, ''))
        cat_id = res[0] if res else None
        if cur_type == ptype and (cat_id is None or cur_cat == cat_id):
            continue
        plan.append((gid, handle, title, ptype, cat_id))
    dist = collections.Counter(p[3] for p in plan)
    print(f'products to update: {len(plan)}; by type: {dist.most_common(8)} ...')
    if a.limit:
        plan = plan[:a.limit]
    if not a.execute:
        print('dry run only; add --execute to write'); return
    ok, errors = 0, []
    mut = 'mutation($p: ProductUpdateInput!) { productUpdate(product: $p) { product { id productType category { id } } userErrors { field message } } }'
    for i, (gid, handle, title, ptype, cat_id) in enumerate(plan, 1):
        p = {'id': gid, 'productType': ptype}
        if cat_id:
            p['category'] = cat_id
        try:
            r = gql(mut, {'p': p})
            ue = r['data']['productUpdate']['userErrors']
            if ue:
                errors.append((handle, title, ue))
            else:
                ok += 1
        except Exception as e:
            errors.append((handle, title, str(e)))
        if i % 50 == 0:
            print(f'{i}/{len(plan)} ok={ok} errors={len(errors)}', flush=True)
        time.sleep(0.25)
    os.makedirs('docs/import/out', exist_ok=True)
    rep = {'updated': ok, 'errors': errors, 'skipped': skipped, 'unresolved_categories': unresolved, 'by_type': dist.most_common()}
    json.dump(rep, open('docs/import/out/product-types-report.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'done: updated {ok}, errors {len(errors)}, skipped {len(skipped)}; report docs/import/out/product-types-report.json')


if __name__ == '__main__':
    main()
